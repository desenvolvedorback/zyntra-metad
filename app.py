# -*- coding: utf-8 -*-
"""
ZYNTRA - Metadata Intelligence
Ferramenta OSINT de extração de metadados de arquivos.

Sem banco de dados, sem IA. Tudo é processado em memória / arquivo
temporário e devolvido como JSON para o frontend.
"""

import base64
import hashlib
import io
import json
import math
import os
import re
import struct
import subprocess
import tempfile
import zipfile
import zlib
from datetime import datetime, timezone

from flask import Flask, jsonify, render_template, request

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 300 * 1024 * 1024  # 300 MB

# ----------------------------------------------------------------------
# Bibliotecas opcionais (import protegido: a ferramenta segue funcionando
# mesmo se alguma faltar no ambiente, só reduz a profundidade da extração)
# ----------------------------------------------------------------------
try:
    from PIL import Image, ExifTags, TiffImagePlugin
    HAS_PIL = True
except Exception:
    HAS_PIL = False

try:
    import pikepdf
    HAS_PIKEPDF = True
except Exception:
    HAS_PIKEPDF = False

try:
    import docx as python_docx
    HAS_DOCX = True
except Exception:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_XLSX = True
except Exception:
    HAS_XLSX = False

try:
    from pptx import Presentation
    HAS_PPTX = True
except Exception:
    HAS_PPTX = False

try:
    from cryptography import x509
    from cryptography.hazmat.primitives.serialization import pkcs7
    from cryptography.hazmat.backends import default_backend
    HAS_CRYPTO = True
except Exception:
    HAS_CRYPTO = False


# ============================================================
# Helpers gerais
# ============================================================

def human_size(n):
    if n is None:
        return None
    step = 1024.0
    units = ["B", "KB", "MB", "GB", "TB"]
    size = float(n)
    for unit in units:
        if size < step:
            return f"{size:.2f} {unit}" if unit != "B" else f"{int(size)} {unit}"
        size /= step
    return f"{size:.2f} PB"


def calc_entropy(data: bytes) -> float:
    if not data:
        return 0.0
    freq = [0] * 256
    for b in data:
        freq[b] += 1
    length = len(data)
    entropy = 0.0
    for f in freq:
        if f:
            p = f / length
            entropy -= p * math.log2(p)
    return round(entropy, 4)


def run_file_cmd(path):
    """Usa o binário `file` (libmagic) para MIME e descrição real."""
    mime, desc = None, None
    try:
        mime = subprocess.run(
            ["file", "--mime-type", "-b", path],
            capture_output=True, text=True, timeout=10
        ).stdout.strip()
    except Exception:
        pass
    try:
        desc = subprocess.run(
            ["file", "-b", path],
            capture_output=True, text=True, timeout=10
        ).stdout.strip()
    except Exception:
        pass
    return mime or "application/octet-stream", desc or "Desconhecido"


def safe_str(v):
    if v is None:
        return None
    if isinstance(v, bytes):
        try:
            return v.decode("utf-8", errors="replace")
        except Exception:
            return repr(v)
    return str(v)


def clean_dict(d):
    """Remove chaves com valor None/vazio para não poluir o JSON final."""
    if not isinstance(d, dict):
        return d
    out = {}
    for k, v in d.items():
        if isinstance(v, dict):
            v = clean_dict(v)
        if v in (None, "", [], {}):
            continue
        out[k] = v
    return out


# ============================================================
# Hashes
# ============================================================

def compute_hashes(data: bytes):
    hashes = {
        "md5": hashlib.md5(data).hexdigest(),
        "sha1": hashlib.sha1(data).hexdigest(),
        "sha224": hashlib.sha224(data).hexdigest(),
        "sha256": hashlib.sha256(data).hexdigest(),
        "sha384": hashlib.sha384(data).hexdigest(),
        "sha512": hashlib.sha512(data).hexdigest(),
        "sha3_256": hashlib.sha3_256(data).hexdigest(),
        "sha3_512": hashlib.sha3_512(data).hexdigest(),
        "blake2b": hashlib.blake2b(data).hexdigest(),
        "blake2s": hashlib.blake2s(data).hexdigest(),
        "crc32": format(zlib.crc32(data) & 0xFFFFFFFF, "08x"),
    }
    return hashes


# ============================================================
# Strings / OSINT patterns
# ============================================================

RE_EMAIL = re.compile(rb"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
RE_URL = re.compile(rb"https?://[^\s\"'<>\\]{4,300}")
RE_IPV4 = re.compile(rb"(?:(?:25[0-5]|2[0-4]\d|1?\d?\d)\.){3}(?:25[0-5]|2[0-4]\d|1?\d?\d)")
RE_IPV6 = re.compile(rb"(?:[0-9a-fA-F]{1,4}:){2,7}[0-9a-fA-F]{1,4}")
RE_DOMAIN = re.compile(rb"\b(?:[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?\.)+(?:com|net|org|io|dev|gov|edu|br|co|app|info|xyz|ai|onion)\b")
RE_JWT = re.compile(rb"eyJ[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}")
RE_UUID = re.compile(rb"[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}")
RE_APIKEY = re.compile(rb"\b(?:AKIA[0-9A-Z]{16}|AIza[0-9A-Za-z_-]{35}|sk-[A-Za-z0-9]{20,}|ghp_[A-Za-z0-9]{30,}|xox[baprs]-[A-Za-z0-9-]{10,})\b")
RE_BASE64 = re.compile(rb"(?:[A-Za-z0-9+/]{4}){10,}(?:[A-Za-z0-9+/]{2}==|[A-Za-z0-9+/]{3}=)?")
RE_USERAGENT = re.compile(rb"Mozilla/5\.0[^\r\n\"'<>]{5,200}")
RE_COOKIE = re.compile(rb"(?:[A-Za-z0-9_]+=[A-Za-z0-9%._-]+; ?){2,}")
RE_PRINTABLE = re.compile(rb"[\x20-\x7e]{5,}")


def extract_strings(data: bytes, limit=60):
    def find(pattern, n=limit):
        found = []
        seen = set()
        for m in pattern.finditer(data):
            s = m.group().decode("latin-1", errors="replace")
            if s not in seen:
                seen.add(s)
                found.append(s)
            if len(found) >= n:
                break
        return found

    result = {
        "emails": find(RE_EMAIL),
        "urls": find(RE_URL),
        "domains": find(RE_DOMAIN, 40),
        "ips_v4": find(RE_IPV4, 40),
        "ips_v6": find(RE_IPV6, 20),
        "jwt_tokens": find(RE_JWT, 10),
        "api_keys": find(RE_APIKEY, 20),
        "uuids": find(RE_UUID, 30),
        "user_agents": find(RE_USERAGENT, 10),
        "cookies": find(RE_COOKIE, 10),
        "printable_sample": find(RE_PRINTABLE, 200),
    }
    return clean_dict(result)


# ============================================================
# Hex viewer
# ============================================================

def hex_dump(data: bytes, offset=0, length=4096):
    chunk = data[offset:offset + length]
    lines = []
    for i in range(0, len(chunk), 16):
        row = chunk[i:i + 16]
        hex_part = " ".join(f"{b:02x}" for b in row)
        ascii_part = "".join(chr(b) if 32 <= b <= 126 else "." for b in row)
        lines.append({
            "offset": f"{offset + i:08x}",
            "hex": hex_part,
            "ascii": ascii_part,
        })
    return {
        "total_size": len(data),
        "shown_bytes": len(chunk),
        "rows": lines,
    }


# ============================================================
# Imagens (EXIF / GPS)
# ============================================================

def _rational_to_float(v):
    try:
        if hasattr(v, "numerator"):
            return float(v.numerator) / float(v.denominator) if v.denominator else 0.0
        if isinstance(v, tuple) and len(v) == 2:
            return float(v[0]) / float(v[1]) if v[1] else 0.0
        return float(v)
    except Exception:
        return None


def _dms_to_decimal(dms, ref):
    try:
        degrees = _rational_to_float(dms[0])
        minutes = _rational_to_float(dms[1])
        seconds = _rational_to_float(dms[2])
        dec = degrees + (minutes / 60.0) + (seconds / 3600.0)
        if ref in ("S", "W"):
            dec = -dec
        return round(dec, 7)
    except Exception:
        return None


def image_info(path):
    if not HAS_PIL:
        return {}, {}
    exif_out, gps_out = {}, {}
    try:
        img = Image.open(path)
        exif_out["formato"] = img.format
        exif_out["modo_cor"] = img.mode
        exif_out["resolucao"] = f"{img.width}x{img.height}"
        exif_out["compressao"] = getattr(img, "compression", None) or img.info.get("compression")
        icc = img.info.get("icc_profile")
        exif_out["perfil_icc"] = "Presente" if icc else None
        if "dpi" in img.info:
            exif_out["dpi"] = str(img.info["dpi"])

        # cor dominante (amostragem rápida)
        try:
            small = img.convert("RGB").resize((50, 50))
            colors = small.getcolors(50 * 50)
            if colors:
                colors.sort(reverse=True, key=lambda c: c[0])
                dom = colors[0][1]
                exif_out["cor_dominante"] = "#{:02x}{:02x}{:02x}".format(*dom)
        except Exception:
            pass

        raw_exif = img.getexif()
        if raw_exif:
            tag_map = {ExifTags.TAGS.get(k, k): v for k, v in raw_exif.items()}

            def g(*names):
                for n in names:
                    if n in tag_map:
                        return tag_map[n]
                return None

            exif_out["fabricante"] = safe_str(g("Make"))
            exif_out["modelo"] = safe_str(g("Model"))
            exif_out["software"] = safe_str(g("Software"))
            exif_out["data_original"] = safe_str(g("DateTimeOriginal", "DateTime"))
            exif_out["orientacao"] = safe_str(g("Orientation"))
            exif_out["exposicao"] = safe_str(g("ExposureTime"))
            exif_out["iso"] = safe_str(g("ISOSpeedRatings", "PhotographicSensitivity"))
            exif_out["abertura"] = safe_str(g("FNumber"))
            exif_out["distancia_focal"] = safe_str(g("FocalLength"))
            exif_out["flash"] = safe_str(g("Flash"))
            exif_out["lente"] = safe_str(g("LensModel"))
            exif_out["copyright"] = safe_str(g("Copyright"))
            exif_out["artista"] = safe_str(g("Artist"))
            exif_out["thumbnail"] = "Presente" if getattr(img, "info", {}).get("thumbnail") else None

            # Histórico de edição (assinatura de softwares comuns)
            sw = safe_str(g("Software")) or ""
            if sw:
                exif_out["historico_edicao"] = f"Última modificação registrada via: {sw}"

            exif_out["assinatura_exif"] = "EXIF presente e legível" if tag_map else None

            # GPS
            gps_ifd = raw_exif.get_ifd(ExifTags.IFD.GPSInfo) if hasattr(ExifTags, "IFD") else None
            if gps_ifd:
                gps_tags = {ExifTags.GPSTAGS.get(k, k): v for k, v in gps_ifd.items()}
                lat = gps_tags.get("GPSLatitude")
                lat_ref = gps_tags.get("GPSLatitudeRef")
                lon = gps_tags.get("GPSLongitude")
                lon_ref = gps_tags.get("GPSLongitudeRef")
                alt = gps_tags.get("GPSAltitude")
                if lat and lon:
                    latitude = _dms_to_decimal(lat, lat_ref)
                    longitude = _dms_to_decimal(lon, lon_ref)
                    if latitude is not None and longitude is not None:
                        gps_out["latitude"] = latitude
                        gps_out["longitude"] = longitude
                        gps_out["altitude"] = _rational_to_float(alt) if alt is not None else None
                        gps_out["google_maps"] = f"https://www.google.com/maps?q={latitude},{longitude}"
                        gps_out["openstreetmap"] = f"https://www.openstreetmap.org/?mlat={latitude}&mlon={longitude}#map=17/{latitude}/{longitude}"
        img.close()
    except Exception as e:
        exif_out["erro"] = f"Falha ao ler EXIF: {e}"
    return clean_dict(exif_out), clean_dict(gps_out)


# ============================================================
# PDF
# ============================================================

def pdf_info(path, raw_data):
    out = {}
    try:
        if HAS_PIKEPDF:
            with pikepdf.open(path) as pdf:
                docinfo = pdf.docinfo or {}
                out["autor"] = safe_str(docinfo.get("/Author"))
                out["empresa"] = safe_str(docinfo.get("/Company"))
                out["criador"] = safe_str(docinfo.get("/Creator"))
                out["produtor_software"] = safe_str(docinfo.get("/Producer"))
                out["data_criacao"] = safe_str(docinfo.get("/CreationDate"))
                out["data_modificacao"] = safe_str(docinfo.get("/ModDate"))
                out["titulo"] = safe_str(docinfo.get("/Title"))
                out["assunto"] = safe_str(docinfo.get("/Subject"))
                out["palavras_chave"] = safe_str(docinfo.get("/Keywords"))
                out["numero_paginas"] = len(pdf.pages)
                out["versao_pdf"] = str(pdf.pdf_version)
                out["criptografado"] = pdf.is_encrypted
                try:
                    xmp = pdf.open_metadata()
                    if xmp:
                        lang = xmp.get("dc:language")
                        out["idioma"] = safe_str(lang)
                        out["metadados_xmp"] = "Presente"
                except Exception:
                    pass
        # Sinais via varredura do conteúdo bruto (funciona mesmo sem pikepdf)
        out["javascript"] = bool(re.search(rb"/JavaScript|/JS\b", raw_data))
        out["campos_ocultos_acroform"] = bool(re.search(rb"/AcroForm", raw_data))
        out["links_uri"] = bool(re.search(rb"/URI\s*\(", raw_data))
        out["assinaturas_digitais"] = bool(re.search(rb"/ByteRange", raw_data))
        out["imagens_internas"] = len(re.findall(rb"/Subtype\s*/Image", raw_data))
        out["objetos_totais"] = len(re.findall(rb"\d+\s+\d+\s+obj", raw_data))
        urls = list(dict.fromkeys(m.group().decode("latin-1", "replace")
                                   for m in RE_URL.finditer(raw_data)))[:30]
        out["urls_encontradas"] = urls
        out["permissoes"] = "Restrita (criptografado)" if out.get("criptografado") else "Sem restrições detectadas"
    except Exception as e:
        out["erro"] = f"Falha ao ler PDF: {e}"
    return clean_dict(out)


# ============================================================
# Office (docx / xlsx / pptx) + detecção de macros
# ============================================================

def office_core_props(cp):
    if cp is None:
        return {}
    return clean_dict({
        "autor": safe_str(getattr(cp, "author", None)),
        "ultimo_editor": safe_str(getattr(cp, "last_modified_by", None)),
        "empresa": safe_str(getattr(cp, "company", None) if hasattr(cp, "company") else None),
        "criado": safe_str(getattr(cp, "created", None)),
        "modificado": safe_str(getattr(cp, "modified", None)),
        "titulo": safe_str(getattr(cp, "title", None)),
        "assunto": safe_str(getattr(cp, "subject", None)),
        "categoria": safe_str(getattr(cp, "category", None)),
        "palavras_chave": safe_str(getattr(cp, "keywords", None)),
        "comentarios": safe_str(getattr(cp, "comments", None)),
        "numero_revisao": safe_str(getattr(cp, "revision", None)),
        "modelo_template": safe_str(getattr(cp, "template", None) if hasattr(cp, "template") else None),
        "aplicativo": safe_str(getattr(cp, "version", None) if hasattr(cp, "version") else None),
    })


def office_info(path, ext):
    out = {"formato": ext.upper()}
    try:
        has_macro = False
        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
                has_macro = any("vbaProject.bin" in n for n in names)
                out["numero_arquivos_internos"] = len(names)
        except Exception:
            pass
        out["macros_detectadas"] = has_macro

        if ext == "docx" and HAS_DOCX:
            doc = python_docx.Document(path)
            out.update(office_core_props(doc.core_properties))
            out["numero_paragrafos"] = len(doc.paragraphs)
            out["numero_tabelas"] = len(doc.tables)
        elif ext in ("xlsx", "xlsm") and HAS_XLSX:
            wb = openpyxl.load_workbook(path, data_only=False, read_only=True)
            out.update(office_core_props(wb.properties))
            out["planilhas"] = wb.sheetnames
            out["numero_planilhas"] = len(wb.sheetnames)
        elif ext == "pptx" and HAS_PPTX:
            pres = Presentation(path)
            out.update(office_core_props(pres.core_properties))
            out["numero_slides"] = len(pres.slides.__iter__.__self__._sldIdLst) if hasattr(pres.slides, "_sldIdLst") else len(list(pres.slides))
    except Exception as e:
        out["erro"] = f"Falha ao ler documento Office: {e}"
    return clean_dict(out)


# ============================================================
# Áudio / Vídeo (via ffprobe)
# ============================================================

def ffprobe_info(path):
    try:
        proc = subprocess.run(
            ["ffprobe", "-v", "quiet", "-print_format", "json",
             "-show_format", "-show_streams", path],
            capture_output=True, text=True, timeout=30
        )
        if proc.returncode != 0 or not proc.stdout:
            return None
        return json.loads(proc.stdout)
    except Exception:
        return None


def media_info(path, kind):
    data = ffprobe_info(path)
    out = {}
    if not data:
        return out
    fmt = data.get("format", {}) or {}
    tags = {k.lower(): v for k, v in (fmt.get("tags") or {}).items()}
    out["duracao_segundos"] = fmt.get("duration")
    out["bitrate"] = fmt.get("bit_rate")
    out["formato_container"] = fmt.get("format_long_name")
    out["artista"] = tags.get("artist")
    out["album"] = tags.get("album")
    out["titulo"] = tags.get("title")
    out["genero"] = tags.get("genre")
    out["ano"] = tags.get("date") or tags.get("year")
    out["encoder"] = tags.get("encoder")
    out["compositor"] = tags.get("composer")
    out["comentario"] = tags.get("comment")
    out["software"] = tags.get("encoding_tool") or tags.get("encoder")
    out["gps"] = tags.get("location") or tags.get("com.apple.quicktime.location.iso6709")
    out["data_criacao"] = tags.get("creation_time")

    streams = data.get("streams", []) or []
    audio_streams = [s for s in streams if s.get("codec_type") == "audio"]
    video_streams = [s for s in streams if s.get("codec_type") == "video"]
    subtitle_streams = [s for s in streams if s.get("codec_type") == "subtitle"]

    if kind == "video" and video_streams:
        v = video_streams[0]
        out["codec_video"] = v.get("codec_long_name") or v.get("codec_name")
        out["largura"] = v.get("width")
        out["altura"] = v.get("height")
        if v.get("width") and v.get("height"):
            out["resolucao"] = f"{v.get('width')}x{v.get('height')}"
        fps = v.get("avg_frame_rate")
        if fps and fps != "0/0":
            try:
                num, den = fps.split("/")
                out["fps"] = round(float(num) / float(den), 2) if float(den) else None
            except Exception:
                out["fps"] = fps
        out["hdr"] = "Sim" if "bt2020" in (v.get("color_space") or "").lower() else "Não detectado"
        out["pixel_format"] = v.get("pix_fmt")

    if audio_streams:
        a = audio_streams[0]
        out["codec_audio"] = a.get("codec_long_name") or a.get("codec_name")
        out["sample_rate"] = a.get("sample_rate")
        out["canais"] = a.get("channels")
        out["bitrate_audio"] = a.get("bit_rate")

    out["legendas"] = len(subtitle_streams) if subtitle_streams else None
    out["miniaturas"] = "Não extraído (requer processamento adicional)" if kind == "video" else None

    return clean_dict(out)


# ============================================================
# ZIP / APK
# ============================================================

def zip_info(path):
    out = {"arquivos_internos": [], "total_arquivos": 0, "tamanho_total_descomprimido": 0}
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            out["total_arquivos"] = len(infos)
            out["comentario_zip"] = safe_str(z.comment) if z.comment else None
            entries = []
            total_uncompressed = 0
            for i in infos[:200]:
                total_uncompressed += i.file_size
                entries.append({
                    "nome": i.filename,
                    "tamanho": i.file_size,
                    "comprimido": i.compress_size,
                    "compressao": i.compress_type,
                    "data": "%04d-%02d-%02d %02d:%02d:%02d" % i.date_time,
                    "crc32": format(i.CRC, "08x"),
                })
            out["arquivos_internos"] = entries
            out["tamanho_total_descomprimido"] = total_uncompressed
    except Exception as e:
        out["erro"] = f"Falha ao ler ZIP: {e}"
    return clean_dict(out)


def apk_info(path):
    out = {}
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            out["manifest_presente"] = "AndroidManifest.xml" in names
            out["dex_files"] = [n for n in names if n.endswith(".dex")]
            out["numero_dex"] = len(out["dex_files"])
            out["bibliotecas_nativas"] = [n for n in names if n.startswith("lib/") and n.endswith(".so")]
            out["assets"] = len([n for n in names if n.startswith("assets/")])
            out["resources_arsc"] = "resources.arsc" in names
            out["arquivos_assinatura"] = [n for n in names if n.startswith("META-INF/") and n.upper().endswith((".RSA", ".DSA", ".EC"))]
            out["manifest_nota"] = ("AndroidManifest.xml está em formato binário AXML; "
                                     "extração de package name/permissions/SDK requer um parser "
                                     "AXML dedicado (não disponível neste ambiente).")
    except Exception as e:
        out["erro"] = f"Falha ao ler APK: {e}"
    return clean_dict(out)


# ============================================================
# Executáveis (PE / ELF / Mach-O) — parsing manual de headers
# ============================================================

def parse_pe(data):
    out = {"tipo": "PE (Windows)"}
    try:
        if data[:2] != b"MZ":
            return None
        e_lfanew = struct.unpack_from("<I", data, 0x3C)[0]
        if data[e_lfanew:e_lfanew + 4] != b"PE\x00\x00":
            return None
        machine, num_sections, timestamp = struct.unpack_from("<HHI", data, e_lfanew + 4)
        machines = {0x14c: "x86 (32-bit)", 0x8664: "x86_64 (64-bit)",
                    0x1c0: "ARM", 0xaa64: "ARM64"}
        out["arquitetura"] = machines.get(machine, hex(machine))
        out["numero_secoes"] = num_sections
        try:
            out["compilado_em"] = datetime.fromtimestamp(timestamp, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        except Exception:
            pass
        opt_hdr_size = struct.unpack_from("<H", data, e_lfanew + 20)[0]
        characteristics = struct.unpack_from("<H", data, e_lfanew + 22)[0]
        out["dll"] = bool(characteristics & 0x2000)
        out["executavel"] = bool(characteristics & 0x0002)
        magic = struct.unpack_from("<H", data, e_lfanew + 24)[0] if opt_hdr_size else None
        out["pe32_plus"] = magic == 0x20b if magic else None
        # Diretório de importação/exportação apenas sinalizados (parsing completo é extenso)
        out["nota"] = "Imports/Exports/DLLs detalhados requerem parsing completo de tabelas (pefile indisponível neste ambiente)."
    except Exception as e:
        out["erro"] = f"Falha ao interpretar PE: {e}"
    return clean_dict(out)


def parse_elf(data):
    out = {"tipo": "ELF (Linux)"}
    try:
        if data[:4] != b"\x7fELF":
            return None
        ei_class = data[4]
        ei_data = data[5]
        out["arquitetura"] = "64-bit" if ei_class == 2 else "32-bit"
        out["endianness"] = "Little-endian" if ei_data == 1 else "Big-endian"
        endian = "<" if ei_data == 1 else ">"
        e_type = struct.unpack_from(endian + "H", data, 16)[0]
        types = {1: "Relocatable", 2: "Executable", 3: "Shared object (.so)", 4: "Core dump"}
        out["tipo_arquivo"] = types.get(e_type, str(e_type))
        e_machine = struct.unpack_from(endian + "H", data, 18)[0]
        machines = {0x3e: "x86_64", 0x28: "ARM", 0xb7: "AArch64", 0x03: "x86"}
        out["maquina"] = machines.get(e_machine, hex(e_machine))
    except Exception as e:
        out["erro"] = f"Falha ao interpretar ELF: {e}"
    return clean_dict(out)


def parse_macho(data):
    out = {"tipo": "Mach-O (macOS)"}
    try:
        magics = {
            b"\xfe\xed\xfa\xce": ("32-bit", "big"),
            b"\xce\xfa\xed\xfe": ("32-bit", "little"),
            b"\xfe\xed\xfa\xcf": ("64-bit", "big"),
            b"\xcf\xfa\xed\xfe": ("64-bit", "little"),
            b"\xca\xfe\xba\xbe": ("fat/universal", None),
        }
        head = bytes(data[:4])
        if head not in magics:
            return None
        arch, endian = magics[head]
        out["arquitetura"] = arch
    except Exception as e:
        out["erro"] = f"Falha ao interpretar Mach-O: {e}"
    return clean_dict(out)


def executable_info(data):
    for parser in (parse_pe, parse_elf, parse_macho):
        result = parser(data)
        if result:
            result["entropia_secoes_nota"] = "Ver entropia geral do arquivo na aba Resumo para indício de packing/ofuscação."
            return result
    return {}


# ============================================================
# Certificados digitais
# ============================================================

def cert_from_x509(cert):
    try:
        return clean_dict({
            "issuer": cert.issuer.rfc4514_string(),
            "subject": cert.subject.rfc4514_string(),
            "serial": format(cert.serial_number, "x"),
            "valido_de": cert.not_valid_before_utc.strftime("%Y-%m-%d %H:%M:%S UTC") if hasattr(cert, "not_valid_before_utc") else str(cert.not_valid_before),
            "valido_ate": cert.not_valid_after_utc.strftime("%Y-%m-%d %H:%M:%S UTC") if hasattr(cert, "not_valid_after_utc") else str(cert.not_valid_after),
            "algoritmo_assinatura": cert.signature_algorithm_oid._name if cert.signature_algorithm_oid else None,
            "fingerprint_sha256": cert.fingerprint(hashlib.sha256()).hex(),
        })
    except Exception:
        return {}


def certificates_info(path, data, ext):
    if not HAS_CRYPTO:
        return {}
    certs_found = []
    try:
        # Arquivo é ele mesmo um certificado?
        if ext in ("pem", "crt", "cer"):
            try:
                certs_found.append(cert_from_x509(x509.load_pem_x509_certificate(data)))
            except Exception:
                try:
                    certs_found.append(cert_from_x509(x509.load_der_x509_certificate(data)))
                except Exception:
                    pass
        elif ext == "der":
            certs_found.append(cert_from_x509(x509.load_der_x509_certificate(data)))

        # APKs assinados: META-INF/*.RSA|*.DSA contém PKCS7 com o certificado
        if ext in ("apk", "jar", "zip"):
            try:
                with zipfile.ZipFile(path) as z:
                    for n in z.namelist():
                        if n.startswith("META-INF/") and n.upper().endswith((".RSA", ".DSA", ".EC")):
                            raw = z.read(n)
                            try:
                                for c in pkcs7.load_der_pkcs7_certificates(raw):
                                    certs_found.append(cert_from_x509(c))
                            except Exception:
                                pass
            except Exception:
                pass
    except Exception:
        pass
    certs_found = [c for c in certs_found if c]
    if not certs_found:
        return {}
    return {"certificados_encontrados": len(certs_found), "certificados": certs_found}


# ============================================================
# Risco
# ============================================================

def assess_risk(ext, mime, data, extra):
    reasons = []
    level = "baixo"

    if extra.get("pdf", {}).get("javascript"):
        reasons.append("PDF contém JavaScript embutido")
        level = "alto"
    if extra.get("office", {}).get("macros_detectadas"):
        reasons.append("Documento Office contém macros (VBA)")
        level = "alto"
    if extra.get("executable"):
        if not extra.get("certificates"):
            reasons.append("Executável sem assinatura digital detectada")
            level = "medio" if level != "alto" else level
    entropy = calc_entropy(data[:2_000_000])
    if entropy > 7.5 and ext in ("exe", "dll", "apk", "bin", ""):
        reasons.append(f"Entropia muito alta ({entropy}) — possível compactação/ofuscação/criptografia")
        level = "medio" if level == "baixo" else level
    if extra.get("strings", {}).get("api_keys"):
        reasons.append("Possíveis chaves de API / tokens expostos no conteúdo")
        level = "alto"
    if extra.get("gps"):
        reasons.append("Arquivo contém coordenadas GPS (dado de geolocalização sensível)")
        level = "medio" if level == "baixo" else level

    if not reasons:
        reasons.append("Nenhum indicador de risco relevante encontrado")

    return {"nivel": level, "motivos": reasons}


# ============================================================
# Orquestrador principal
# ============================================================

IMAGE_EXT = {"jpg", "jpeg", "png", "webp", "heic", "heif", "tif", "tiff", "bmp", "gif", "raw", "cr2", "nef", "arw"}
AUDIO_EXT = {"mp3", "flac", "aac", "wav", "ogg", "m4a", "wma"}
VIDEO_EXT = {"mp4", "mkv", "avi", "mov", "webm", "flv", "wmv"}
OFFICE_EXT = {"docx", "xlsx", "xlsm", "pptx", "docm", "dotx", "odt", "ods", "odp"}


def analyze_file(path, original_name):
    with open(path, "rb") as f:
        data = f.read()

    ext = (original_name.rsplit(".", 1)[-1].lower() if "." in original_name else "")
    mime, description = run_file_cmd(path)
    size = len(data)

    result = {}

    result["hashes"] = compute_hashes(data)

    exif, gps = ({}, {})
    if HAS_PIL and (ext in IMAGE_EXT or mime.startswith("image/")):
        exif, gps = image_info(path)
    if exif:
        result["exif"] = exif
    if gps:
        result["gps"] = gps

    if ext == "pdf" or mime == "application/pdf":
        result["pdf"] = pdf_info(path, data)

    if ext in OFFICE_EXT or "officedocument" in mime or "ms-office" in mime or "opendocument" in mime:
        office_ext = ext if ext in ("docx", "xlsx", "xlsm", "pptx") else (
            "docx" if "wordprocessingml" in mime else
            "xlsx" if "spreadsheetml" in mime else
            "pptx" if "presentationml" in mime else ext
        )
        result["office"] = office_info(path, office_ext or "office")

    if ext in AUDIO_EXT or mime.startswith("audio/"):
        m = media_info(path, "audio")
        if m:
            result["audio"] = m

    if ext in VIDEO_EXT or mime.startswith("video/"):
        m = media_info(path, "video")
        if m:
            result["video"] = m

    if ext == "apk":
        result["zip"] = zip_info(path)
        result["apk"] = apk_info(path)
    elif ext == "zip" or mime in ("application/zip", "application/x-zip-compressed"):
        result["zip"] = zip_info(path)

    if ext in ("exe", "dll", "so", "bin", "elf", "") or mime in (
        "application/x-dosexec", "application/x-executable",
        "application/x-sharedlib", "application/x-elf-executable",
        "application/x-mach-binary",
    ):
        exe = executable_info(data)
        if exe:
            result["executable"] = exe

    certs = certificates_info(path, data, ext)
    if certs:
        result["certificates"] = certs

    result["strings"] = extract_strings(data)
    result["hex"] = hex_dump(data, 0, 4096)

    result["risk"] = assess_risk(ext, mime, data, result)

    entropy = calc_entropy(data if size < 5_000_000 else data[:5_000_000])

    result["file"] = clean_dict({
        "nome": original_name,
        "tipo_mime": mime,
        "descricao_real": description,
        "extensao": ext or "(sem extensão)",
        "tamanho_bytes": size,
        "tamanho_legivel": human_size(size),
        "entropia": entropy,
        "sha256": result["hashes"]["sha256"],
        "md5": result["hashes"]["md5"],
        "analisado_em": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
    })

    # Ordena para o formato pedido: file, hashes, depois o resto
    ordered = {"file": result["file"], "hashes": result["hashes"], "risk": result["risk"]}
    for k in ("exif", "gps", "pdf", "office", "audio", "video", "zip", "apk",
              "executable", "certificates", "strings", "hex"):
        if k in result:
            ordered[k] = result[k]
    return ordered


# ============================================================
# Rotas
# ============================================================

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/analyze", methods=["POST"])
def api_analyze():
    if "file" not in request.files:
        return jsonify({"erro": "Nenhum arquivo enviado."}), 400
    f = request.files["file"]
    if not f.filename:
        return jsonify({"erro": "Nome de arquivo inválido."}), 400

    tmp_dir = tempfile.mkdtemp(prefix="zyntra_")
    tmp_path = os.path.join(tmp_dir, "upload")
    try:
        f.save(tmp_path)
        result = analyze_file(tmp_path, f.filename)
        return jsonify(result)
    except Exception as e:
        return jsonify({"erro": f"Falha ao analisar arquivo: {e}"}), 500
    finally:
        try:
            os.remove(tmp_path)
            os.rmdir(tmp_dir)
        except Exception:
            pass


@app.errorhandler(413)
def too_large(e):
    return jsonify({"erro": "Arquivo excede o limite máximo de 300MB."}), 413


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
